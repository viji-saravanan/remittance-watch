"""Streaming parser for the RPW workbook's current-era sheet.

Ground truth: ``docs/research/01-rpw-dataset.md``. The 48 MB workbook's main
sheet uncompresses to ~298 MB, so everything here is ``read_only`` streaming —
no random access, no materialized lists.

Contract enforced at parse time (structural drift raises loudly, never silently):
* the sheet must exist under its exact name;
* the header must equal EXPECTED_COLUMNS exactly (42 named columns);
* rows may carry junk beyond column 42 (observed: stray values in 43–47) —
  ignored, and excluded from the idempotency hash;

The decomposition identity ``total% = fee/lcu_amount×100 + fx_margin`` holds for
99.98% of transparent tiers, but ~85 tiers per full release (all 2016–2020 waves)
violate it upstream. Those are recorded as MathViolations (published totals are
kept verbatim — they remain the WB's official numbers), not raised — UNLESS no
MathAudit is supplied (strict mode), which raises FormatDrift on the first one.
"""

from collections.abc import Iterator
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .normalize import (
    normalize_instrument,
    parse_bool_transparent,
    parse_decimal,
    parse_period,
    row_hash,
)

SHEET_NAME = "Dataset (from Q2 2016)"

EXPECTED_COLUMNS = (
    "id", "period",
    "source_code", "source_name", "source_region", "source_income", "source_lending", "source_G8G20",
    "destination_code", "destination_name", "destination_region", "destination_income", "destination_lending", "destination_G8G20",
    "firm", "firm_type", "payment instrument", "access point", "speed actual",
    "cc1 lcu amount", "cc1 denomination amount", "cc1 lcu code", "cc1 lcu fee", "cc1 lcu fx rate", "cc1 fx margin", "cc1 total cost %",
    "cc2 lcu amount", "cc2 denomination amount", "cc2 lcu code", "cc2 lcu fee", "cc2 lcu fx rate", "cc2 fx margin", "cc2 total cost %",
    "inter lcu bank fx", "transparent", "Standard Note", "note2",
    "receiving network coverage", "pickup location", "pickup method", "date", "corridor",
)

# Decomposition identity tolerance, in percentage points (source rounds to 2dp).
MATH_TOLERANCE = Decimal("0.05")

# Column indexes within EXPECTED_COLUMNS — the parser's only access pattern.
COL = {name: idx for idx, name in enumerate(EXPECTED_COLUMNS)}


class FormatDrift(Exception):
    """The workbook no longer matches the contract — fail loudly, fix deliberately."""


@dataclass(slots=True)
class MathViolation:
    row_number: int
    source_iso3: str
    destination_iso3: str
    tier: str
    expected_pct: Decimal
    published_pct: Decimal


@dataclass(slots=True)
class MathAudit:
    """Collects decomposition-identity violations instead of raising.

    ``checked`` counts every transparent, fully-populated tier the identity was
    tested on, so the caller can report a true violation rate (full-release
    baseline: 85 / 401,156 ≈ 0.021%, all in 2016–2020 waves).
    """

    checked: int = 0
    violations: list[MathViolation] = field(default_factory=list)
    max_examples: int = 10

    @property
    def violation_count(self) -> int:
        return len(self.violations)

    def violation_rate(self) -> Decimal:
        if self.checked == 0:
            return Decimal(0)
        return Decimal(self.violation_count) / Decimal(self.checked) * 100

    def report(self) -> str:
        head = (
            f"math audit: {self.violation_count} identity violations "
            f"across {self.checked:,} checked tiers ({self.violation_rate():.4f}%)"
        )
        return head


@dataclass(slots=True)
class CountryRef:
    iso3: str
    name: str
    region: str | None
    income: str | None
    lending: str | None
    g8g20: str | None


@dataclass(slots=True)
class TierBlock:
    lcu_amount: Decimal | None
    denomination_usd: Decimal | None
    currency: str | None
    fee_lcu: Decimal | None
    fx_rate: Decimal | None
    fx_margin_pct: Decimal | None
    total_cost_pct: Decimal | None

    def is_empty(self) -> bool:
        return (
            self.lcu_amount is None
            and self.fee_lcu is None
            and self.total_cost_pct is None
        )


@dataclass(slots=True)
class SourceRow:
    period: int
    source: CountryRef
    destination: CountryRef
    firm: str
    firm_type: str | None
    instrument: str
    access_point: str | None
    speed_actual: str | None
    interbank_fx_rate: Decimal | None
    transparent: bool
    note: str | None
    cc1: TierBlock
    cc2: TierBlock
    raw_row_hash: str
    row_number: int

    def tiers(self) -> Iterator[tuple[str, TierBlock]]:
        yield "cc1", self.cc1
        yield "cc2", self.cc2


def parse_workbook(path: str | Path, *, audit: MathAudit | None = None) -> Iterator[SourceRow]:
    """Yield SourceRows from the current-era sheet, streaming.

    Pass ``audit`` to record identity violations and continue; omit it (or pass
    None) for strict mode, where the first violation raises FormatDrift.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            msg = f"sheet {SHEET_NAME!r} not found; sheets present: {wb.sheetnames}"
            raise FormatDrift(msg)
        ws = wb[SHEET_NAME]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            msg = f"sheet {SHEET_NAME!r} is empty"
            raise FormatDrift(msg)
        named = header[: len(EXPECTED_COLUMNS)]
        if tuple(named) != EXPECTED_COLUMNS:
            expected_set, got_set = set(EXPECTED_COLUMNS), set(named)
            missing = expected_set - got_set
            added = got_set - expected_set
            msg = (
                "workbook header drifted from the contract. "
                f"missing: {sorted(missing)}; unexpected: {sorted(added)}; "
                f"order changed: {tuple(named) != EXPECTED_COLUMNS and not missing and not added}"
            )
            raise FormatDrift(msg)

        for row_number, cells in enumerate(rows, start=2):  # header consumed; real sheets are 1-based
            if not cells:
                continue
            # Normalize width both ways: producers trim trailing blanks (rows come
            # back short) and the real file reports junk beyond column 42.
            canonical = (tuple(cells) + (None,) * len(EXPECTED_COLUMNS))[: len(EXPECTED_COLUMNS)]
            if not any(c is not None and str(c) != "" for c in canonical):
                continue
            yield _parse_row(canonical, row_number, audit)
    finally:
        wb.close()


def _parse_row(cells: tuple, row_number: int, audit: MathAudit | None) -> SourceRow:
    def text(name: str) -> str | None:
        value = cells[COL[name]]
        if value is None:
            return None
        stripped = str(value).strip()
        return stripped or None

    def dec(name: str) -> Decimal | None:
        return parse_decimal(cells[COL[name]])

    source = CountryRef(
        iso3=_require(cells[COL["source_code"]], "source_code", row_number),
        name=_require(cells[COL["source_name"]], "source_name", row_number),
        region=text("source_region"),
        income=text("source_income"),
        lending=text("source_lending"),
        g8g20=text("source_G8G20"),
    )
    destination = CountryRef(
        iso3=_require(cells[COL["destination_code"]], "destination_code", row_number),
        name=_require(cells[COL["destination_name"]], "destination_name", row_number),
        region=text("destination_region"),
        income=text("destination_income"),
        lending=text("destination_lending"),
        g8g20=text("destination_G8G20"),
    )

    row = SourceRow(
        period=parse_period(cells[COL["period"]]),
        source=source,
        destination=destination,
        firm=_require(cells[COL["firm"]], "firm", row_number),
        firm_type=text("firm_type"),
        instrument=normalize_instrument(cells[COL["payment instrument"]]),
        access_point=text("access point"),
        speed_actual=text("speed actual"),
        interbank_fx_rate=dec("inter lcu bank fx"),
        transparent=parse_bool_transparent(cells[COL["transparent"]]),
        note=text("Standard Note"),
        cc1=_tier(cells, "cc1"),
        cc2=_tier(cells, "cc2"),
        raw_row_hash=row_hash(cells),
        row_number=row_number,
    )
    _check_math(row, audit)
    return row


def _tier(cells: tuple, prefix: str) -> TierBlock:
    return TierBlock(
        lcu_amount=parse_decimal(cells[COL[f"{prefix} lcu amount"]]),
        denomination_usd=parse_decimal(cells[COL[f"{prefix} denomination amount"]]),
        currency=cells[COL[f"{prefix} lcu code"]] and str(cells[COL[f"{prefix} lcu code"]]).strip() or None,
        fee_lcu=parse_decimal(cells[COL[f"{prefix} lcu fee"]]),
        fx_rate=parse_decimal(cells[COL[f"{prefix} lcu fx rate"]]),
        fx_margin_pct=parse_decimal(cells[COL[f"{prefix} fx margin"]]),
        total_cost_pct=parse_decimal(cells[COL[f"{prefix} total cost %"]]),
    )


def _check_math(row: SourceRow, audit: MathAudit | None) -> None:
    """Transparent rows should satisfy the decomposition identity.

    Violations are upstream data-quality quirks (85 per full release, 2016–2020
    only): recorded when auditing, raised when strict.
    """
    for tier_name, tier in row.tiers():
        if not row.transparent or tier.is_empty():
            continue
        if None in (tier.lcu_amount, tier.fee_lcu, tier.fx_margin_pct, tier.total_cost_pct):
            continue  # partial blocks exist upstream; loader stores what's there
        if tier.lcu_amount == 0:
            continue  # identity is undefined on a zero base — store verbatim, don't divide
        if audit is not None:
            audit.checked += 1
        fee_pct = tier.fee_lcu / tier.lcu_amount * 100
        expected = fee_pct + tier.fx_margin_pct
        if abs(expected - tier.total_cost_pct) > MATH_TOLERANCE:
            violation = MathViolation(
                row_number=row.row_number,
                source_iso3=row.source.iso3,
                destination_iso3=row.destination.iso3,
                tier=tier_name,
                expected_pct=expected,
                published_pct=tier.total_cost_pct,
            )
            if audit is None:
                msg = (
                    f"row {row.row_number} ({row.source.iso3}->{row.destination.iso3}, "
                    f"{tier_name}): {tier.fee_lcu}/{tier.lcu_amount}*100 + "
                    f"{tier.fx_margin_pct} = {expected} but total cost reads "
                    f"{tier.total_cost_pct} — upstream math or our parsing drifted"
                )
                raise FormatDrift(msg)
            audit.violations.append(violation)


def _require(value: object, column: str, row_number: int) -> str:
    text = str(value).strip() if value is not None else ""
    if not text:
        msg = f"row {row_number}: required column {column!r} is empty"
        raise FormatDrift(msg)
    return text
