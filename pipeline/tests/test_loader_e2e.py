"""End-to-end ingest against a real Postgres: counts, acceptance SQL,
idempotency, and the non-transparent NULL-margin rule. Skips cleanly when
no Postgres is reachable (see conftest)."""

import pytest

from remittance_watch.loader import ingest_rows
from remittance_watch.parser import parse_workbook

pytestmark = pytest.mark.usefixtures("clean_db")


@pytest.fixture()
def do_ingest(migrated_conn, fixture_path):
    def run():
        return ingest_rows(migrated_conn, parse_workbook(fixture_path))

    return run


class TestIngest:
    def test_counts(self, do_ingest):
        stats = do_ingest()
        assert stats.source_rows == 4
        assert stats.quote_rows == 8  # every row has both tiers populated
        assert stats.skipped_tiers == 0
        assert stats.quarters == {20253, 20251}

    def test_acceptance_sql_decomposition(self, migrated_conn, do_ingest):
        """The milestone question: corridor X, amount Y -> fee vs margin per provider."""
        from decimal import Decimal

        do_ingest()
        row = migrated_conn.execute(
            """
            SELECT p.name, qu.amount_usd, qu.fee_lcu / qu.lcu_amount * 100 AS fee_pct,
                   qu.fx_margin_pct, qu.total_cost_pct
            FROM quotes qu
            JOIN quarters q   ON q.id = qu.quarter_id
            JOIN corridors c  ON c.id = qu.corridor_id
            JOIN providers p  ON p.id = qu.provider_id
            WHERE c.code = 'AGONAM' AND qu.amount_usd = 200 AND qu.transparent
              AND p.name = 'Example MTO Ltd'
            """
        ).fetchone()
        assert row is not None
        name, amount, fee_pct, margin_pct, total_pct = row
        assert name == "Example MTO Ltd"
        assert amount == Decimal("200")
        # NUMERIC comes back as Decimal — compare exactly, never through float.
        assert fee_pct == Decimal("2257.2") / Decimal("33000") * 100
        assert margin_pct == Decimal("3.10")
        assert total_pct == Decimal("9.94")
        # decomposition identity holds in SQL too
        assert fee_pct + margin_pct - total_pct <= Decimal("0.05")

    def test_non_transparent_margin_is_null(self, migrated_conn, do_ingest):
        """transparent='no' stores a lying literal 0 upstream — we store NULL + flag."""
        do_ingest()
        rows = migrated_conn.execute(
            """
            SELECT qu.fx_margin_pct, qu.total_cost_pct, qu.transparent
            FROM quotes qu
            JOIN providers p ON p.id = qu.provider_id
            WHERE p.name = 'Example Bank NA'
            """
        ).fetchall()
        assert len(rows) == 2
        for margin, total, transparent in rows:
            assert margin is None
            assert transparent is False
            assert total is not None  # WB's published number kept, flagged non-transparent

    def test_idempotent_reingest(self, migrated_conn, do_ingest):
        do_ingest()
        before = migrated_conn.execute("SELECT COUNT(*) FROM quotes").fetchone()[0]
        stats = do_ingest()
        after = migrated_conn.execute("SELECT COUNT(*) FROM quotes").fetchone()[0]
        assert after == before == 8
        assert stats.quote_rows == 8  # upserted, not duplicated

    def test_correction_updates_in_place(self, migrated_conn, fixture_path, tmp_path):
        """A revised workbook (changed fee) updates the existing quote row."""
        from openpyxl import Workbook, load_workbook

        from remittance_watch.parser import EXPECTED_COLUMNS, SHEET_NAME

        src = load_workbook(fixture_path, read_only=True)
        data = [tuple(r) for r in src[SHEET_NAME].iter_rows(values_only=True)]
        src.close()
        col = list(EXPECTED_COLUMNS).index("cc1 lcu fee")
        data[1] = tuple(data[1][:col]) + (2500.0,) + tuple(data[1][col + 1:])
        col_total = list(EXPECTED_COLUMNS).index("cc1 total cost %")
        data[1] = tuple(data[1][:col_total]) + (10.68,) + tuple(data[1][col_total + 1:])  # 2500/33000*100+3.10

        out = tmp_path / "revised.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        for record in data:
            ws.append(record)
        wb.save(out)

        ingest_rows(migrated_conn, parse_workbook(fixture_path))
        ingest_rows(migrated_conn, parse_workbook(out))
        fee = migrated_conn.execute(
            """
            SELECT qu.fee_lcu FROM quotes qu
            JOIN providers p ON p.id = qu.provider_id
            JOIN corridors c ON c.id = qu.corridor_id
            WHERE p.name = 'Example MTO Ltd' AND c.code = 'AGONAM' AND qu.amount_usd = 200
            """
        ).fetchone()[0]
        assert float(fee) == 2500.0
        assert migrated_conn.execute("SELECT COUNT(*) FROM quotes").fetchone()[0] == 8
