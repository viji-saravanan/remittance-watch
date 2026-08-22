"""Contract tests: the parser against the committed fixture workbook.

The fixture (tests/make_fixture.py) replicates the real sheet's structure.
Upstream format drift must fail LOUDLY here, not silently in prod
(milestone acceptance).
"""

import shutil

import pytest
from openpyxl import Workbook, load_workbook

from remittance_watch.parser import (
    EXPECTED_COLUMNS,
    SHEET_NAME,
    FormatDrift,
    parse_workbook,
)


def _rows(path):
    return list(parse_workbook(path))


def _mutated_copy(fixture_path, tmp_path, mutate_header=None, mutate_cell=None):
    """Rebuild the workbook with a tampered header/cell; returns new path."""
    src = load_workbook(fixture_path, read_only=True)
    data = [tuple(r) for r in src[SHEET_NAME].iter_rows(values_only=True)]
    src.close()

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    if mutate_header:
        data[0] = mutate_header(list(data[0]))
    if mutate_cell:
        data = mutate_cell(data)
    for record in data:
        ws.append(record)
    out = tmp_path / "mutated.xlsx"
    wb.save(out)
    return out


class TestHappyPath:
    def test_parses_all_fixture_rows(self, fixture_path):
        rows = _rows(fixture_path)
        assert len(rows) == 4

    def test_row1_fields(self, fixture_path):
        from decimal import Decimal

        row = _rows(fixture_path)[0]
        assert row.period == 20253
        assert row.source.iso3 == "AGO" and row.destination.iso3 == "NAM"
        assert row.firm == "Example MTO Ltd"
        assert row.instrument == "Cash"
        assert row.transparent is True
        assert row.interbank_fx_rate == Decimal("160")
        # Decimals compare exactly against strings — never float literals.
        assert row.cc1.fee_lcu == Decimal("2257.2")
        assert row.cc1.total_cost_pct == Decimal("9.94")

    def test_instrument_combo_normalized(self, fixture_path):
        row = next(r for r in _rows(fixture_path) if r.row_number == 4)  # 'Cash,mobile MONEY'
        assert row.instrument == "Cash, Mobile money"

    def test_hashes_stable_across_reads(self, fixture_path):
        hashes = [r.raw_row_hash for r in _rows(fixture_path)]
        assert hashes == [r.raw_row_hash for r in _rows(fixture_path)]
        assert len(set(hashes)) == 4


class TestMathContract:
    def test_transparent_rows_satisfy_identity(self, fixture_path):
        from decimal import Decimal

        for row in _rows(fixture_path):
            if not row.transparent:
                continue
            for tier in (row.cc1, row.cc2):
                if tier.is_empty() or None in (tier.lcu_amount, tier.fee_lcu, tier.fx_margin_pct, tier.total_cost_pct):
                    continue
                expected = tier.fee_lcu / tier.lcu_amount * 100 + tier.fx_margin_pct
                assert abs(expected - tier.total_cost_pct) <= Decimal("0.05")


class TestLoudDrift:
    def test_renamed_column_raises(self, fixture_path, tmp_path):
        path = _mutated_copy(
            fixture_path, tmp_path,
            mutate_header=lambda h: ["cc1 cost %" if c == "cc1 total cost %" else c for c in h],
        )
        with pytest.raises(FormatDrift, match="cc1 cost %"):
            list(parse_workbook(path))

    def test_missing_column_raises(self, fixture_path, tmp_path):
        path = _mutated_copy(
            fixture_path, tmp_path,
            mutate_header=lambda h: h[:20] + h[21:],  # drop one column
        )
        with pytest.raises(FormatDrift):
            list(parse_workbook(path))

    def test_tampered_total_cost_raises(self, fixture_path, tmp_path):
        # cc1 total cost % lives at column index of "cc1 total cost %"; corrupt a data row.
        col = list(EXPECTED_COLUMNS).index("cc1 total cost %")

        def corrupt(data):
            data[1] = tuple(data[1][:col]) + (99.99,) + tuple(data[1][col + 1:])
            return data

        path = _mutated_copy(fixture_path, tmp_path, mutate_cell=corrupt)
        with pytest.raises(FormatDrift, match="drifted"):
            list(parse_workbook(path))

    def test_missing_sheet_raises(self, fixture_path, tmp_path):
        out = tmp_path / "nosheet.xlsx"
        shutil.copy(fixture_path, out)
        wb = load_workbook(out)
        # openpyxl refuses to save a workbook with zero visible sheets, so
        # leave a placeholder behind when removing the contract sheet.
        wb.create_sheet("unrelated")
        del wb[SHEET_NAME]
        wb.save(out)
        with pytest.raises(FormatDrift, match="sheet"):
            list(parse_workbook(out))


class TestMathAuditMode:
    """Real releases carry ~85 identity violations (2016–2020 waves). Audited
    parsing records them and keeps the published numbers; strict raises."""

    def test_violation_recorded_and_row_kept(self, fixture_path, tmp_path):
        from decimal import Decimal

        from remittance_watch.parser import MathAudit

        col = list(EXPECTED_COLUMNS).index("cc1 total cost %")

        def corrupt(data):
            data[1] = tuple(data[1][:col]) + (99.99,) + tuple(data[1][col + 1:])
            return data

        path = _mutated_copy(fixture_path, tmp_path, mutate_cell=corrupt)
        audit = MathAudit()
        rows = list(parse_workbook(path, audit=audit))

        assert len(rows) == 4  # nothing dropped
        assert audit.checked >= 6  # rows 1/3/4 are transparent with full tiers
        assert audit.violation_count == 1
        v = audit.violations[0]
        assert v.tier == "cc1" and v.row_number == 2
        assert str(v.published_pct) == "99.99"
        # the published (wrong) number is preserved verbatim — it's still upstream's
        assert rows[0].cc1.total_cost_pct == Decimal("99.99")

    def test_clean_workbook_audits_zero(self, fixture_path):
        from remittance_watch.parser import MathAudit

        audit = MathAudit()
        list(parse_workbook(fixture_path, audit=audit))
        assert audit.violation_count == 0
        assert audit.checked > 0
